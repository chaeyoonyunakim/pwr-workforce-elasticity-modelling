"""Source reader contract tests.

Each reader is exercised against a small fixture file built in ``tmp_path``
to keep tests independent of the (gitignored) ``/data/`` directory. Tests
pin the schema, the in-window filter behaviour and at least one numeric
expectation per reader.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from pwr_elasticity import io

# ---------------------------------------------------------------------------
# Module-level smoke check on the public surface
# ---------------------------------------------------------------------------


def test_module_exposes_expected_readers() -> None:
    expected = {
        "read_tac",
        "read_hchs_staff_in_post",
        "read_hchs_turnover",
        "read_vacancies",
        "read_earnings",
        "read_ae",
        "read_rtt",
        "read_ods_trusts",
    }
    assert expected.issubset(dir(io))


# ---------------------------------------------------------------------------
# read_ods_trusts
# ---------------------------------------------------------------------------


def _write_ods_csv(path: Path) -> None:
    rows = [
        # Currently open trust
        ["R1A"] + ["NORTH TRUST", "Y56", "QKK"] + [""] * 6 + ["20100401", ""] + [""] * 15,
        # Closed trust
        ["R1B"] + ["LEGACY TRUST", "Y62", "QXX"] + [""] * 6 + ["20100401", "20200331"] + [""] * 15,
    ]
    lines = [",".join(f'"{c}"' for c in row) for row in rows]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def test_read_ods_trusts_schema_and_open_flag(tmp_path: Path) -> None:
    csv = tmp_path / "etr.csv"
    _write_ods_csv(csv)
    df = io.read_ods_trusts(csv)
    assert list(df.columns) == [
        "org_code",
        "org_name",
        "nhse_region_code",
        "parent_org_code",
        "open_date",
        "close_date",
        "is_open",
    ]
    assert len(df) == 2
    assert (
        df.loc[df["org_code"] == "R1A", "is_open"].iloc[0] is True
        or df.loc[df["org_code"] == "R1A", "is_open"].iloc[0]
    )
    assert not df.loc[df["org_code"] == "R1B", "is_open"].iloc[0]
    assert df["open_date"].dt.tz is not None


# ---------------------------------------------------------------------------
# read_earnings
# ---------------------------------------------------------------------------


def test_read_earnings_filters_to_window(tmp_path: Path) -> None:
    csv = tmp_path / "earnings.csv"
    pd.DataFrame(
        {
            "DATE": ["2019-04-30", "2021-04-30", "2024-10-31"],
            "STAFF_GROUP_ORDER": [1, 1, 1],
            "STAFF_GROUP": ["All staff", "All staff", "All staff"],
            "PAYMENT_TYPE": ["A", "A", "A"],
            "AMOUNT": [100.0, 200.0, 300.0],
            "SAMPLE_SIZE": [10, 20, 30],
        }
    ).to_csv(csv, index=False)
    df = io.read_earnings(csv)
    # Pre-window row 2019-04 dropped; two in-window rows remain.
    assert set(df["mean_basic_pay_gbp"]) == {200.0, 300.0}
    assert df["period_month"].dt.tz is not None
    assert df["afc_band"].iloc[0] == "All AfC bands"


# ---------------------------------------------------------------------------
# read_hchs_staff_in_post
# ---------------------------------------------------------------------------


def test_read_hchs_staff_in_post_keeps_org_level(tmp_path: Path) -> None:
    csv = tmp_path / "hchs.csv"
    pd.DataFrame(
        {
            "DATA_MONTH": ["2020-09-30", "2021-09-30", "2024-09-30"],
            "DATA_LEVEL": ["Organisation", "National", "Organisation"],
            "NHSE_REGION_CODE": ["Y56", "England", "Y56"],
            "NHSE_REGION_NAME": ["London", "England", "London"],
            "ICS_CODE": ["QWE", "", "QWE"],
            "ICS_NAME": ["A", "", "A"],
            "ORG_CODE": ["R1A", "ALL", "R1A"],
            "ORG_NAME": ["North", "ALL", "North"],
            "CLUSTER_GROUP": ["", "", ""],
            "BENCHMARK_GROUP": ["", "", ""],
            "MAIN_STAFF_GROUP": ["All", "All", "All"],
            "STAFF_GROUP_1": ["All", "All", "Doctors"],
            "FTE": [1000.0, 1_000_000.0, 1100.0],
            "HEADCOUNT": [1100, 1_200_000, 1200],
        }
    ).to_csv(csv, index=False)
    df = io.read_hchs_staff_in_post(csv)
    # National-aggregate row dropped (DATA_LEVEL filter); pre-window row dropped.
    assert len(df) == 1
    assert df["org_code"].iloc[0] == "R1A"
    assert df["staff_group"].iloc[0] == "Doctors"


# ---------------------------------------------------------------------------
# read_hchs_turnover
# ---------------------------------------------------------------------------


def test_read_hchs_turnover_filters_pre_2022(tmp_path: Path) -> None:
    csv = tmp_path / "turnover.csv"
    pd.DataFrame(
        {
            "PERIOD": ["201009 to 201109", "202109 to 202209", "202309 to 202409"],
            "TYPE": ["Denominator", "Denominator", "Denominator"],
            "NHSE_REGION_CODE": ["Y56", "Y56", "Y56"],
            "NHSE_REGION_NAME": ["London", "London", "London"],
            "MAIN_STAFF_GROUP": ["All", "All", "All"],
            "STAFF_GROUP": ["All", "All", "All"],
            "HC": [1000, 1100, 1200],
            "FTE": [950.0, 1050.0, 1150.0],
        }
    ).to_csv(csv, index=False)
    df = io.read_hchs_turnover(csv)
    assert set(df["period_end_year"].tolist()) == {2022, 2024}


# ---------------------------------------------------------------------------
# read_vacancies
# ---------------------------------------------------------------------------


def _write_vacancy_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Total 2018 onwards"
    ws.append(["NHS Vacancy Statistics"])
    for _ in range(18):
        ws.append([])
    ws.append(["Total workforce vacancy FTE"])
    ws.append(["Region", "Sector", "2021/22 Q1 (Jun-21)", "2024/25 Q3 (Dec-24)"])
    ws.append(["London", "Acute", 1000, 1200])
    ws.append(["London Total", None, 1000, 1200])
    ws.append([])
    ws.append(["Total workforce % vacancy rate"])
    ws.append(["Region", "Sector", "2021/22 Q1 (Jun-21)", "2024/25 Q3 (Dec-24)"])
    ws.append(["London", "Acute", 8.5, 10.2])
    ws.append(["London Total", None, 8.5, 10.2])
    wb.save(path)


def test_read_vacancies_joins_fte_and_rate(tmp_path: Path) -> None:
    xlsx = tmp_path / "vac.xlsx"
    _write_vacancy_xlsx(xlsx)
    df = io.read_vacancies(xlsx)
    assert set(df.columns) == {
        "nhse_region_name",
        "sector",
        "period_quarter",
        "vacancy_fte",
        "vacancy_rate",
    }
    assert len(df) == 2
    # Rate normalised from percent to fraction.
    assert df["vacancy_rate"].max() < 1.0
    # Aggregate 'Total' rows must be filtered out.
    assert not df["sector"].astype(str).str.contains("Total").any()


# ---------------------------------------------------------------------------
# read_ae
# ---------------------------------------------------------------------------


def _write_ae_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Activity"
    ws.append(["A&E Time Series"])
    for _ in range(11):
        ws.append([])
    # Header row 12: super-header
    ws.append([None, None, "A&E attendances", None, None, "Emergency Admissions"])
    # Header row 13: column names (matches reader's regex)
    ws.append(
        [
            None,
            "Period",
            "Type 1 Departments - Major A&E",
            "Type 3 Departments - Other A&E/Minor Injury Unit",
            "Total Attendances",
            "Total Emergency Admissions",
            "Number of patients spending >4 hours from decision to admit to admission",
            "Number of patients spending >12 hours from decision to admit to admission",
        ]
    )
    # Data rows
    ws.append([None, pd.Timestamp("2020-09-01"), 900_000, 400_000, 1_400_000, 300_000, 50_000, 100])
    ws.append(
        [None, pd.Timestamp("2021-09-01"), 1_000_000, 500_000, 1_600_000, 320_000, 60_000, 200]
    )
    ws.append(
        [None, pd.Timestamp("2024-10-01"), 1_100_000, 550_000, 1_700_000, 340_000, 80_000, 250]
    )
    wb.save(path)


def test_read_ae_window_and_pct(tmp_path: Path) -> None:
    xlsx = tmp_path / "ae.xlsx"
    _write_ae_xlsx(xlsx)
    df = io.read_ae(xlsx)
    # Pre-window 2020-09 dropped.
    assert len(df) == 2
    # pct_met_4hr = 1 - (over_4hr / total_attendances). For 2021-09: 1 - 60000/1600000 = 0.9625
    row = df[df["period_month"].dt.year == 2021].iloc[0]
    assert abs(row["pct_met_4hr"] - (1 - 60_000 / 1_600_000)) < 1e-9


# ---------------------------------------------------------------------------
# read_rtt
# ---------------------------------------------------------------------------


def _write_rtt_csv(path: Path) -> None:
    # Minimal columns: Period, RTT Part Type, Treatment Function Code/Name,
    # Provider Org Code, and at least one weekly bucket plus the >52 / >104
    # buckets needed by the reader.
    cols = [
        "Period",
        "Provider Org Code",
        "Provider Org Name",
        "RTT Part Type",
        "Treatment Function Code",
        "Treatment Function Name",
        "Gt 00 To 01 Weeks SUM 1",
        "Gt 51 To 52 Weeks SUM 1",
        "Gt 52 To 53 Weeks SUM 1",
        "Gt 104 Weeks SUM 1",
    ]
    rows = [
        # In-scope: Part_2, non-C_999 treatment function
        ["RTT-February-2026", "RXX", "TRUST X", "Part_2", "C_101", "Urology", 100, 5, 3, 1],
        # Out-of-scope: C_999 totals row (must be excluded)
        ["RTT-February-2026", "RXX", "TRUST X", "Part_2", "C_999", "Total", 200, 10, 6, 2],
        # Out-of-scope: Part_1 (Completed) — must be excluded
        ["RTT-February-2026", "RXX", "TRUST X", "Part_1", "C_101", "Urology", 50, 0, 0, 0],
    ]
    pd.DataFrame(rows, columns=cols).to_csv(path, index=False)


def test_read_rtt_filters_and_aggregates(tmp_path: Path) -> None:
    csv = tmp_path / "rtt.csv"
    _write_rtt_csv(csv)
    df = io.read_rtt(csv)
    assert len(df) == 1
    row = df.iloc[0]
    assert row["org_code"] == "RXX"
    assert row["treatment_function_code"] == "C_101"
    assert int(row["n_waiting"]) == 100 + 5 + 3 + 1
    assert int(row["n_waiting_52plus_weeks"]) == 3 + 1


# ---------------------------------------------------------------------------
# read_tac
# ---------------------------------------------------------------------------


def _write_tac_xlsx(path: Path) -> None:
    wb = Workbook()
    # Sheet 1: List of providers (case matches 2024/25 vintage)
    ws_list = wb.active
    ws_list.title = "List of providers"
    ws_list.append(["Full name of Provider", "NHS code", "Region", "Sector", "Comments"])
    ws_list.append(["Test Acute Trust", "RXX", "London", "Acute", None])
    # Sheet 2: All data
    ws_data = wb.create_sheet("All data")
    ws_data.append(
        [
            "OrganisationName",
            "WorkSheetName",
            "TableID",
            "MainCode",
            "RowNumber",
            "SubCode",
            "Total",
        ]
    )
    # TAC09 Staff, SubCode STA0366 (Net employee benefits expenditure) for one org
    ws_data.append(["Test Acute Trust", "TAC09 Staff", 2, "A09CY01", 66, "STA0366", 200_000])
    ws_data.append(["Test Acute Trust", "TAC09 Staff", 2, "A09CY01P", 66, "STA0366", 180_000])
    ws_data.append(["Test Acute Trust", "TAC09 Staff", 2, "A09CY01O", 66, "STA0366", 20_000])
    # Noise: different subcode that should be ignored
    ws_data.append(["Test Acute Trust", "TAC09 Staff", 2, "A09CY01", 49, "STA0230", 999_999])
    wb.save(path)


def test_read_tac_extracts_three_pay_categories(tmp_path: Path) -> None:
    fp = tmp_path / "TAC-data-published-in-NHS-trusts-accounts-for-2022-23.xlsx"
    _write_tac_xlsx(fp)
    df = io.read_tac(tmp_path)
    assert len(df) == 1
    row = df.iloc[0]
    assert row["org_code"] == "RXX"
    assert row["financial_year"] == "2022/23"
    # Values are stored in £'000 in TAC; reader scales to £.
    assert row["substantive_pay_gbp"] == 180_000 * 1_000
    assert row["other_staff_pay_gbp"] == 20_000 * 1_000
    assert row["total_pay_gbp"] == 200_000 * 1_000


def test_read_tac_skips_out_of_window_files(tmp_path: Path) -> None:
    fp = tmp_path / "TAC-data-published-in-NHS-trusts-accounts-for-2019-20.xlsx"
    _write_tac_xlsx(fp)
    df = io.read_tac(tmp_path)
    assert df.empty


# ---------------------------------------------------------------------------
# Cross-cutting: every reader returns a DataFrame (smoke)
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "fn",
    [
        io.read_tac,
        io.read_hchs_staff_in_post,
        io.read_hchs_turnover,
        io.read_vacancies,
        io.read_earnings,
        io.read_ae,
        io.read_rtt,
        io.read_ods_trusts,
    ],
)
def test_reader_signature_is_callable(fn) -> None:
    assert callable(fn)
